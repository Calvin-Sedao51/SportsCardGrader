"""Panini America adapter: sealed products from the public sitemap; checklists only
when a file URL is already known.

What answers and what does not (recon 2026-10-17, browser UA, from a US IP):

* ``https://www.paniniamerica.net/`` and every ``.html`` page, ``/checklist.html``,
  ``/graphql``, ``/rest/V1/...`` and the ``blog.``/``college.``/``nft.`` hosts: HTTP
  403 with a 730 KB Cloudflare custom block page (``::CLOUDFLARE_ERROR_1000S_BOX::``
  placeholder, a base64 PNG on a dark background) for curl. httpx gets the 14 KB
  SPA shell instead, which carries no product data: the storefront is a React app
  that loads everything through ``POST /onepanini`` (GraphQL).
* ``POST /onepanini``: 403 for plain HTTP whatever the headers. In a browser each
  request carries a ``signature`` header computed by the SPA's own JS and the API
  answers 426 "Invalid request" without it, so the API is not replayable outside
  the page and we do not try to reverse that.
* Playwright: the bundled headless shell gets the document but 403 on the API
  (``sec-ch-ua`` advertises HeadlessChrome). The full Chromium build in new-headless
  mode did get listing and product JSON for a few pages, but after a handful of
  clicks on ``/checklist.html`` the site replaced the page with "We have detected
  suspicious activity ... you are blocked from proceeding further ... refrain from
  any unethical practices". That is an explicit no, so no browser path ships.
* ``robots.txt`` (200; only ``/cdn-cgi/``, ``/upload/invoice/``, ``/upload/tax/`` and
  ``/support/*.pdf`` disallowed) advertises ``sitemap.xml`` (200, ~4,200 URLs). The
  sealed-product slugs in it are the one machine-readable listing that answers, so
  releases come from there. Names are rebuilt from the slug (flagged in ``raw``);
  release dates, images and checklist links are not in the sitemap.
* Files under ``https://assets.paniniamerica.net/checklist/`` download fine (200)
  when the URL is known, but the index that maps products to those xlsx/pdf files
  lives behind the blocked API, so ``fetch_checklist`` can only parse a release
  that already carries a ``checklist_url`` and returns ``[]`` otherwise (the run
  stays ``ok``: the listing worked; cards simply aren't obtainable yet).
"""
import re
from io import BytesIO
from typing import Optional
from urllib.parse import urlsplit

from openpyxl import load_workbook

from app.sportsdb.http import Blocked
from app.sportsdb.models import Card, Release, guess_sport, guess_year

SITEMAP_URL = "https://www.paniniamerica.net/sitemap.xml"
STORE_HOST = "www.paniniamerica.net"   # college. and nft. hosts list NIL singles and NFTs
CHECKLIST_BLOCKED = ("paniniamerica.net: the checklist index is only served through the storefront's "
                     "signed GraphQL API, which 403s plain HTTP and blocks headless browsers; "
                     "see module docstring")

_URL_ENTRY = re.compile(r"<url>(.*?)</url>", re.S)
_LOC = re.compile(r"<loc>\s*([^<\s]+)\s*</loc>")
_LASTMOD = re.compile(r"<lastmod>\s*(\d{4}-\d{2}-\d{2})")
_SEASON = re.compile(r"^((?:19|20)\d{2})-(\d{2})-")
# Sealed products start with the year and name a box or pack; singles start with the
# player, Panini Instant with "pi-", stickers are not trading cards.
_PRODUCT = re.compile(r"^(?:19|20)\d{2}.*?(?:-box(?:-|$)|-pack$|blaster|mega-box|breakaway)")
_BOX_VARIANT = re.compile(r"-box-(.+)$")
_ACRONYMS = {"nba", "wnba", "nfl", "mlb", "nhl", "nwsl", "fifa", "efl", "mlw", "nil", "liv", "ufc",
             "pfl", "usa", "mma", "wwe"}
# guess_sport only knows the league keywords; Panini's European basketball brand is not one.
_SPORT_HINTS = {"euroleague": "basketball"}

_HEADER_ALIASES = {
    "card_number": ("card #", "card#", "card number", "card no", "no.", "no", "#", "number"),
    "player": ("player", "athlete", "name", "player name", "description"),
    "team": ("team",),
    "subset": ("set", "set name", "subset", "insert", "program"),
    "parallel": ("parallel", "version", "variation"),
    "print_run": ("seq", "sequence", "print run", "numbered", "numbered to", "qty", "#'d", "serial"),
}
_ROOKIE = re.compile(r"\b(rookie|rookies|rc)\b", re.I)
_AUTO = re.compile(r"\b(auto|autos|autograph|autographs|signature|signatures|signed)\b", re.I)
_RELIC = re.compile(r"\b(relic|relics|memorabilia|material|materials|patch|patches|jersey|jerseys|swatch)\b", re.I)


class PaniniSource:
    name = "panini"

    async def fetch_releases(self, fetcher) -> list[Release]:
        return parse_sitemap(await fetcher.get_text(SITEMAP_URL))

    async def fetch_checklist(self, fetcher, release: Release) -> list[Card]:
        if not release.checklist_url:
            return []    # see CHECKLIST_BLOCKED: no index maps products to files
        if not urlsplit(release.checklist_url).path.lower().endswith(".xlsx"):
            return []
        return parse_checklist_xlsx(await fetcher.get_bytes(release.checklist_url))


def parse_sitemap(xml: str) -> list[Release]:
    releases = []
    for entry in _URL_ENTRY.findall(xml):
        loc = _LOC.search(entry)
        if not loc:
            continue
        parts = urlsplit(loc.group(1))
        slug, dot, ext = parts.path.strip("/").rpartition(".")
        if parts.netloc != STORE_HOST or "/" in slug or ext != "html" or "sticker" in slug:
            continue
        if not _PRODUCT.search(slug):
            continue
        lastmod = _LASTMOD.search(entry)
        name = _name_from_slug(slug)
        releases.append(Release(
            manufacturer="panini", source_id=slug, name=name, sport=_sport(slug, name),
            year=guess_year(name), product_type=_product_type(slug), url=loc.group(1),
            raw={"slug": slug, "lastmod": lastmod.group(1) if lastmod else None,
                 "name_from": "sitemap slug"}))
    return releases


def _name_from_slug(slug: str) -> str:
    """'2025-26-panini-select-basketball-trading-card-box-hobby' ->
    '2025-26 Panini Select Basketball Trading Card Box (Hobby)' (Panini's own pattern)."""
    season = _SEASON.match(slug)
    head = f"{season.group(1)}-{season.group(2)}" if season else ""
    rest = slug[season.end():] if season else slug
    variant = _BOX_VARIANT.search(rest)
    if variant:
        rest = rest[:variant.start()] + "-box"
    words = [head] if head else []
    words += [_word(w) for w in rest.split("-") if w]
    name = " ".join(words)
    return f"{name} ({_title(variant.group(1))})" if variant else name


def _word(token: str) -> str:
    token = re.sub(r"^(\d{4})tm$", r"\1", token)   # '2026tm' is '2026™' in the slug
    return token.upper() if token in _ACRONYMS else token.capitalize()


def _title(tokens: str) -> str:
    return " ".join(_word(t) for t in tokens.split("-") if t)


def _product_type(slug: str) -> Optional[str]:
    variant = _BOX_VARIANT.search(slug)
    if variant:
        return _title(variant.group(1))
    if slug.endswith("-pack"):
        return "Pack"
    return None


def _sport(slug: str, name: str):
    for hint, sport in _SPORT_HINTS.items():
        if hint in slug:
            return sport
    return guess_sport(name)


def parse_checklist_xlsx(data: bytes) -> list[Card]:
    """Header-driven: the header row is wherever a 'card #' and a 'player' column first
    appear together (Panini sheets often start with a title row); other columns optional."""
    cards: list[Card] = []
    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        columns = None
        for row in rows:
            cells = ["" if v is None else str(v).strip() for v in row]
            if columns is None:
                columns = _columns(cells)
                if columns is not None:
                    headers = cells
                continue
            number = cells[columns["card_number"]] if columns["card_number"] < len(cells) else ""
            player = cells[columns["player"]] if columns["player"] < len(cells) else ""
            if not number or not player:
                continue
            get = lambda key: (cells[columns[key]] or None) if key in columns and columns[key] < len(cells) else None
            subset, parallel = get("subset"), get("parallel")
            flags_text = " ".join(filter(None, (subset, parallel)))
            cards.append(Card(
                card_number=number, player=player, team=get("team"), subset=subset, parallel=parallel,
                print_run=_print_run(get("print_run")),
                is_rookie=bool(_ROOKIE.search(flags_text)), is_auto=bool(_AUTO.search(flags_text)),
                is_relic=bool(_RELIC.search(flags_text)),
                raw={h: c for h, c in zip(headers, cells) if h and c}))
    wb.close()
    return cards


def _columns(cells: list[str]) -> Optional[dict[str, int]]:
    lowered = [c.lower() for c in cells]
    found = {}
    for key, aliases in _HEADER_ALIASES.items():
        for i, cell in enumerate(lowered):
            if cell in aliases and i not in found.values():
                found[key] = i
                break
    return found if "card_number" in found and "player" in found else None


def _print_run(value: Optional[str]) -> Optional[int]:
    """'/99', '99', '1 of 1', '#'d/25' -> 99, 99, 1, 25."""
    if not value:
        return None
    if re.search(r"\b1\s*(?:of|/)\s*1\b", value):
        return 1
    m = re.search(r"(\d+)\s*$", value)
    return int(m.group(1)) if m else None
